#!/usr/bin/env python3
"""Append chronological V3, V4, V5, ... judge stages to the RAG workbook.

Run from repository root (Windows PowerShell), for example:
    python -m scripts.import_judge_results --workbook ... --evaluation ... \
        --version v5 --answers ... --questions ... --output ... --dry-run

Dependencies: openpyxl (no artifact_tool needed).

Existing historical versions, scores, formulas, and disagreement notes are
not changed. V4+ requires the preceding version to be present.
Supports additional question IDs with blank human scores. Output must be new.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

HUMAN = "Human Evaluation V2"
JUDGE = "LLM Evaluation V3"
COMPARISON = "Evaluation Comparison"
SUMMARY = "Updated Summary"
CRITERIA = (
    # label, JSON key, Human V2 score column, V3 judge score column, reason column
    ("Correctness", "correctness", "I", "D", "H"),
    ("Completeness", "completeness", "K", "E", "I"),
    ("Citation Support", "citation_support", "M", "F", "J"),
    ("Abstention/Limits", "abstention_limits", "O", "G", "K"),
)
# Existing comparison layout: A:AM is historical and MUST remain in place.
# Append four V3 triplets (Human V2 reused, LLM V3, V3 Agree?) starting AN.
V3_START = 40  # AN
V3_MATCHES = 52       # AZ
V3_COMPARABLE = 53    # BA
V3_AGREEMENT = 54     # BB
V3_NOTES = 55         # BC
NAVY = "17365D"
LIGHT = "EAF2FA"


def load_index(path: str | None, label: str, *, required=False) -> dict:
    if not path:
        if required:
            raise ValueError(f"--{label} is required")
        return {}
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"{label} file not found: {source}")
    data = json.loads(source.read_text(encoding="utf-8-sig"))
    if not isinstance(data, list) or (required and not data):
        raise ValueError(f"{label} must contain a nonempty JSON array")
    result = {}
    for item in data:
        if not isinstance(item, dict):
            raise ValueError(f"{label} must contain JSON objects")
        qid = item.get("question_id")
        if not isinstance(qid, str) or not qid.strip():
            raise ValueError(f"{label}: missing question_id")
        qid = qid.strip()
        if qid in result:
            raise ValueError(f"{label}: duplicate question_id {qid}")
        result[qid] = item
    return result


def validate_evaluations(records):
    for qid, item in records.items():
        if item.get("error"):
            raise ValueError(f"{qid}: evaluation error: {item['error']}")
        scores = item.get("scores")
        if not isinstance(scores, dict) or not isinstance(scores.get("reasons"), dict):
            raise ValueError(f"{qid}: missing scores or reasons")
        for label, key, *_ in CRITERIA:
            if key not in scores:
                raise ValueError(f"{qid}: missing {key}")
            value = scores[key]
            if value is None and key in ("completeness", "citation_support"):
                pass
            elif type(value) is not int or value not in (0, 1, 2):
                raise ValueError(f"{qid}: {key} must be 0, 1, 2 (or null for optional criteria)")
            reason = scores["reasons"].get(key)
            if not isinstance(reason, str) or not reason.strip():
                raise ValueError(f"{qid}: missing {label} reason")


def id_rows(ws):
    result = {}
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if not isinstance(value, str) or not re.fullmatch(r"qst_[A-Za-z0-9_-]+", value.strip()):
            continue
        qid = value.strip()
        if qid in result:
            raise ValueError(f"Duplicate {qid} in {ws.title}")
        result[qid] = row
    return result


def comparison_rows(ws, wb):
    result = {}
    # Older comparison sheets link A to Human Evaluation!A2 with a formula.
    pattern = re.compile(r"'(Human Evaluation(?: V2)?)'!\$?A\$?(\d+)")
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if not isinstance(value, str):
            continue
        if re.fullmatch(r"qst_[A-Za-z0-9_-]+", value.strip()):
            qid = value.strip()
        elif value.startswith("="):
            m = pattern.search(value)
            if not m:
                continue
            qid = wb[m.group(1)].cell(int(m.group(2)), 1).value
            if not isinstance(qid, str) or not qid.startswith("qst_"):
                continue
        else:
            continue
        if qid in result:
            raise ValueError(f"Duplicate comparison row for {qid}")
        result[qid] = row
    return result


def style_header(cell):
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def expected_headers():
    result = {}
    for idx, (label, *_rest) in enumerate(CRITERIA):
        start = V3_START + idx * 3
        result[start] = f"{label} Human V2 (reused for V3)"
        result[start + 1] = f"{label} LLM V3"
        result[start + 2] = f"{label} V3 Agree?"
    result.update({
        V3_MATCHES: "V3 Matches",
        V3_COMPARABLE: "V3 Comparable",
        V3_AGREEMENT: "V3 Agreement %",
        V3_NOTES: "V3 Disagreement Notes",
    })
    return result


def ensure_v3_headers(ws):
    for col, title in expected_headers().items():
        old = ws.cell(1, col).value
        if old not in (None, title):
            raise ValueError(
                f"{COMPARISON}!{get_column_letter(col)}1 already contains {old!r}. "
                "Start from your original rag_evaluation_results_updated.xlsx, "
                "not a workbook produced by the older V3 importer."
            )
        ws.cell(1, col).value = title
        style_header(ws.cell(1, col))
        ws.column_dimensions[get_column_letter(col)].width = 27 if col != V3_NOTES else 55
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 0, 48)


def ensure_judge_sheet(wb):
    headers = (
        "Question ID", "Category", "Generation Status", "Correctness",
        "Completeness", "Citation Support", "Abstention/Limits",
        "Correctness Reason", "Completeness Reason", "Citation Reason",
        "Abstention Reason", "Raw Judge Record",
    )
    if JUDGE in wb:
        ws = wb[JUDGE]
        for col, title in enumerate(headers, 1):
            if ws.cell(1, col).value != title:
                raise ValueError(f"Unexpected existing {JUDGE} header in column {col}")
        return ws
    ws = wb.create_sheet(JUDGE, wb.sheetnames.index("LLM Evaluation V2") + 1)
    ws.append(headers)
    for cell in ws[1]:
        style_header(cell)
    ws.row_dimensions[1].height = 38
    ws.freeze_panes = "D2"
    for col, width in {"A": 19, "B": 24, "C": 23, "D": 18, "E": 18,
                       "F": 18, "G": 22, "H": 55, "I": 55, "J": 55,
                       "K": 55, "L": 75}.items():
        ws.column_dimensions[col].width = width
    return ws


def metadata(qid, answers, questions, record):
    a = answers.get(qid, {})
    q = questions.get(qid, {})
    return (
        a.get("question_type") or q.get("question_type") or "",
        q.get("question") or a.get("question") or "",
        q.get("gold_answer") or q.get("reference_answer") or "",
        q.get("answer_facts") or [],
        a.get("generated_answer") or "",
        a.get("status") or record.get("generation_status") or "",
        a.get("citation_ids") or [],
    )


def as_text(value):
    return "\n".join(map(str, value)) if isinstance(value, list) else str(value or "")


def safe_fill(ws, row, col, value, qid):
    """Only write new/empty V3 cells; refuse conflicting imported records."""
    if value in (None, ""):
        return
    old = ws.cell(row, col).value
    if old in (None, ""):
        ws.cell(row, col).value = value
    elif old != value:
        raise ValueError(f"Conflict: {qid} at {ws.title}!{get_column_letter(col)}{row}; "
                         "existing value differs. No output saved.")


def numeric_or_na_link(sheet, column, row):
    if row is None:
        return '=""'
    ref = f"'{sheet}'!{column}{row}"
    return f'=IF(ISNUMBER({ref}),{ref},IF({ref}="N/A","N/A",""))'


def formula_for_match(h, j):
    # Identical to the existing V1/V2 Match/Disagree convention:
    # two explicit N/A entries match; missing values are not compared.
    return f'=IF(OR({h}="",{j}=""),"",IF({h}={j},"Match","Disagree"))'


def extend_comparison(wb, human_rows, judge_rows):
    ws = wb[COMPARISON]
    ensure_v3_headers(ws)
    rows = comparison_rows(ws, wb)
    new_count = 0
    for qid, human_row in human_rows.items():
        row = rows.get(qid)
        if row is None:
            row = max(ws.max_row + 1, 2)
            rows[qid] = row
            new_count += 1
            # New questions have no V1/V2 historical judge scores.
            # A:G metadata is linked to the new Human V2 row.
            for col in range(1, 8):
                ref = f"'{HUMAN}'!{get_column_letter(col)}{human_row}"
                ws.cell(row, col).value = f'=IF({ref}="","",{ref})'
            # Human V2 source columns in historical comparison (K,Q,W,AC).
            # These are NEW cells in a NEW row; existing rows are untouched.
            for (_label, _key, hcol, *_), old_comp_col in zip(
                CRITERIA, ("K", "Q", "W", "AC")
            ):
                ws[f"{old_comp_col}{row}"] = numeric_or_na_link(HUMAN, hcol, human_row)
        for idx, (_label, _key, hcol, jcol, _rcol) in enumerate(CRITERIA):
            base = V3_START + idx * 3
            h_cell = f"{get_column_letter(base)}{row}"
            j_cell = f"{get_column_letter(base+1)}{row}"
            m_cell = f"{get_column_letter(base+2)}{row}"
            ws.cell(row, base).value = numeric_or_na_link(HUMAN, hcol, human_row)
            ws.cell(row, base + 1).value = numeric_or_na_link(JUDGE, jcol, judge_rows.get(qid))
            ws.cell(row, base + 2).value = formula_for_match(h_cell, j_cell)
        matches = [f"{get_column_letter(V3_START+i*3+2)}{row}" for i in range(4)]
        ws.cell(row, V3_MATCHES).value = "=" + "+".join(f'COUNTIF({m},"Match")' for m in matches)
        ws.cell(row, V3_COMPARABLE).value = "=" + "+".join(
            f'COUNTIF({m},"Match")+COUNTIF({m},"Disagree")' for m in matches
        )
        ws.cell(row, V3_AGREEMENT).value = f'=IF(BA{row}=0,"",AZ{row}/BA{row})'
        ws.cell(row, V3_AGREEMENT).number_format = "0.0%"
        # BC is a manually editable V3 notes column. Never write over it.
    ws.auto_filter.ref = f"A1:BC{max(ws.max_row, 2)}"
    return new_count


def find_summary_cell(ws, row, title):
    """Append a new header without replacing existing summary columns."""
    for col in range(1, ws.max_column + 1):
        if ws.cell(row, col).value == title:
            return col
    col = ws.max_column + 1
    ws.cell(row, col).value = title
    style_header(ws.cell(row, col))
    ws.column_dimensions[get_column_letter(col)].width = 30
    return col


def extend_summary(wb, human_rows, judge_rows):
    ws = wb[SUMMARY]
    last_comp = max(wb[COMPARISON].max_row, 2)
    last_judge = max(wb[JUDGE].max_row, 2)
    # Append V3 rubric averages and sample counts after existing V1/V2 columns.
    avg_col = find_summary_cell(ws, 1, "LLM V3 Avg Rubric Score")
    count_col = find_summary_cell(ws, 1, "LLM V3 n")
    for row, (_label, _key, _hcol, judge_col, _rcol) in enumerate(CRITERIA, 2):
        ref = f"'{JUDGE}'!{judge_col}2:{judge_col}{last_judge}"
        cell = ws.cell(row, avg_col)
        cell.value = f'=IF(COUNT({ref})=0,"",AVERAGE({ref})/2)'
        cell.number_format = "0.0%"
        ws.cell(row, count_col).value = f'=COUNT({ref})'

    # Historical V1/V2 rows 8/9 are left completely unchanged.
    # V3 row 10 is the natural next chronological row in the uploaded master.
    if ws["A10"].value not in (None, "V3 (LLM vs Human V2 reused)"):
        raise ValueError("Updated Summary!A10 is occupied; refusing to overwrite historical content")
    ws["A10"] = "V3 (LLM vs Human V2 reused)"
    ws["B10"] = f"=SUM('{COMPARISON}'!AZ2:AZ{last_comp})"
    ws["C10"] = f"=SUM('{COMPARISON}'!BA2:BA{last_comp})"
    ws["D10"] = '=IF(C10=0,"",B10/C10)'
    ws["D10"].number_format = "0.0%"
    for cell in ws[10][:4]:
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        cell.font = Font(bold=True)

    # If MAE is already in the summary (user's master), extend that existing
    # chronological table with a V3 column; do not add per-question error columns.
    mae_header_row = next((r for r in range(1, ws.max_row + 1)
                           if ws.cell(r, 1).value == "Criterion"
                           and ws.cell(r, 2).value == "V1 MAE (0–2)"), None)
    if mae_header_row is not None:
        mae_col = next((c for c in range(1, ws.max_column + 1)
                        if ws.cell(mae_header_row, c).value == "V3 MAE (0–2)"), None)
        if mae_col is None:
            mae_col = max((c for c in range(1, ws.max_column + 1)
                           if ws.cell(mae_header_row, c).value is not None), default=1) + 1
            # Never overwrite existing values beneath an unlabeled column.
            while any(ws.cell(mae_header_row + i, mae_col).value is not None for i in range(5)):
                mae_col += 1
        ws.cell(mae_header_row, mae_col).value = "V3 MAE (0–2)"
        style_header(ws.cell(mae_header_row, mae_col))
        ws.column_dimensions[get_column_letter(mae_col)].width = 22
        for idx, (_label, _key, hcol, jcol, _reason) in enumerate(CRITERIA, 1):
            r = mae_header_row + idx
            h = f"'{HUMAN}'!{hcol}2:{hcol}{max(wb[HUMAN].max_row,2)}"
            # MAE must align by question ID, not by raw judge sheet row.
            # V3 comparison columns are joined by question ID.
            hc = f"'{COMPARISON}'!{get_column_letter(V3_START+(idx-1)*3)}2:{get_column_letter(V3_START+(idx-1)*3)}{last_comp}"
            jc = f"'{COMPARISON}'!{get_column_letter(V3_START+(idx-1)*3+1)}2:{get_column_letter(V3_START+(idx-1)*3+1)}{last_comp}"
            ws.cell(r, mae_col).value = (
                f'=IFERROR(SUMPRODUCT(IFERROR(ABS({hc}-{jc}),0))/'
                f'SUMPRODUCT(--ISNUMBER({hc}),--ISNUMBER({jc})),"")'
            )
            ws.cell(r, mae_col).number_format = "0.00"
    else:
        # For older workbooks without an MAE section, add one BELOW existing data.
        r0 = ws.max_row + 2
        ws.cell(r0, 1).value = "V3 MAE vs Human V2 (0–2, numeric pairs only)"
        style_header(ws.cell(r0, 1))
        for idx, (label, _key, _hcol, _jcol, _reason) in enumerate(CRITERIA):
            row = r0 + 1 + idx
            hc = f"'{COMPARISON}'!{get_column_letter(V3_START+idx*3)}2:{get_column_letter(V3_START+idx*3)}{last_comp}"
            jc = f"'{COMPARISON}'!{get_column_letter(V3_START+idx*3+1)}2:{get_column_letter(V3_START+idx*3+1)}{last_comp}"
            ws.cell(row, 1).value = label
            ws.cell(row, 2).value = (
                f'=IFERROR(SUMPRODUCT(IFERROR(ABS({hc}-{jc}),0))/'
                f'SUMPRODUCT(--ISNUMBER({hc}),--ISNUMBER({jc})),"")'
            )
            ws.cell(row, 2).number_format = "0.00"
    # Do not add common-cohort comparisons or rewrite historical metrics.


def run(args):
    src = Path(args.workbook)
    dst = Path(args.output)
    if not src.is_file():
        raise FileNotFoundError(f"Input workbook not found: {src}")
    if src.resolve() == dst.resolve():
        raise ValueError("--output must be a DIFFERENT filename from --workbook")
    if dst.exists() and not args.dry_run:
        raise FileExistsError(f"Output already exists: {dst}. Choose another filename.")
    if args.version.lower() != "v3":
        raise ValueError("This importer supports --version v3")
    evaluations = load_index(args.evaluation, "evaluation", required=True)
    validate_evaluations(evaluations)
    answers = load_index(args.answers, "answers")
    questions = load_index(args.questions, "questions")
    wb = load_workbook(src)
    for sheet in (HUMAN, "LLM Evaluation", "LLM Evaluation V2", COMPARISON, SUMMARY):
        if sheet not in wb:
            raise ValueError(f"Missing expected sheet: {sheet}")
    # Check for conflicting V3 headers even during dry run.
    summary_slot = wb[SUMMARY]["A10"].value
    if summary_slot not in (None, "V3 (LLM vs Human V2 reused)"):
        raise ValueError("Updated Summary!A10 is occupied; no historical content will be overwritten")
    for col, expected in expected_headers().items():
        existing = wb[COMPARISON].cell(1, col).value
        if existing not in (None, expected):
            raise ValueError(f"{COMPARISON}!{get_column_letter(col)}1 already used by {existing!r}")
    human_rows = id_rows(wb[HUMAN])
    comp_rows = comparison_rows(wb[COMPARISON], wb)
    if not human_rows or not comp_rows:
        raise ValueError("Expected existing human and comparison question rows")
    judge_rows = id_rows(wb[JUDGE]) if JUDGE in wb else {}
    # Validate existing judge sheet headers before dry-run success.
    if JUDGE in wb:
        ensure_judge_sheet(wb)
    # Reject missing metadata for genuinely new questions; no placeholder/fabrication.
    new_questions = [qid for qid in evaluations if qid not in human_rows]
    for qid in new_questions:
        m = metadata(qid, answers, questions, evaluations[qid])
        if not m[1]:
            raise ValueError(f"New question {qid} lacks question metadata. "
                             "Supply --answers and --questions containing that ID.")
    # Preflight existing V3 scores/reasons before saving any changes.
    if JUDGE in wb:
        judge = wb[JUDGE]
        for qid, record in evaluations.items():
            if qid not in judge_rows:
                continue
            row = judge_rows[qid]
            expected = {}
            for idx, (_label, key, *_rest) in enumerate(CRITERIA):
                expected[4 + idx] = "N/A" if record["scores"][key] is None else record["scores"][key]
                expected[8 + idx] = record["scores"]["reasons"][key]
            expected[12] = json.dumps(record, ensure_ascii=False)
            for col, value in expected.items():
                old = judge.cell(row, col).value
                if old not in (None, "", value):
                    raise ValueError(f"Conflicting V3 score for {qid} at {get_column_letter(col)}{row}")
    new_judge = [qid for qid in evaluations if qid not in judge_rows]
    new_comparison = [qid for qid in evaluations if qid not in comp_rows]
    if args.dry_run:
        print(f"DRY RUN: V3 records={len(evaluations)}, new judge rows={len(new_judge)}, "
              f"new questions={len(new_questions)}, new comparison rows={len(new_comparison)}")
        print("New question IDs: " + (", ".join(new_questions) if new_questions else "(none)"))
        print("Would add: LLM Evaluation V3; V3 triplets in Evaluation Comparison (AN:AY); "
              "V3 totals (AZ:BC); V3 metrics in Updated Summary")
        print("No files written; original human/V1/V2 data and historical columns untouched.")
        return
    human = wb[HUMAN]
    judge = ensure_judge_sheet(wb)
    next_h = human.max_row + 1
    next_j = judge.max_row + 1
    for qid, record in evaluations.items():
        meta = metadata(qid, answers, questions, record)
        if qid not in human_rows:
            hr = next_h
            next_h += 1
            human_rows[qid] = hr
            values = (qid, meta[0], meta[1], meta[2], as_text(meta[3]), meta[4], meta[5], as_text(meta[6]))
            for col, value in enumerate(values, 1):
                human.cell(hr, col).value = value
            # Scores I/K/M/O and reasons J/L/N/P remain BLANK for independent review.
            human.cell(hr, 17).value = "New question — Human V2 review pending"
        if qid not in judge_rows:
            judge_rows[qid] = next_j
            next_j += 1
        jr = judge_rows[qid]
        for col, value in enumerate((qid, meta[0], meta[5]), 1):
            safe_fill(judge, jr, col, value, qid)
        for idx, (_label, key, *_rest) in enumerate(CRITERIA):
            score = record["scores"][key]
            safe_fill(judge, jr, 4 + idx, "N/A" if score is None else score, qid)
            safe_fill(judge, jr, 8 + idx, record["scores"]["reasons"][key], qid)
        safe_fill(judge, jr, 12, json.dumps(record, ensure_ascii=False), qid)
    judge.auto_filter.ref = f"A1:L{judge.max_row}"
    extend_comparison(wb, human_rows, judge_rows)
    extend_summary(wb, human_rows, judge_rows)
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    print(f"Saved: {dst}")
    print(f"Imported V3 records={len(evaluations)}, new judge rows={len(new_judge)}, "
          f"new questions={len(new_questions)}, new comparison rows={len(new_comparison)}")
    print("Original V1/V2/human scores and historical comparison structure preserved.")


# ------------------- Version-agnostic chronological extension -------------------
# V3 occupies AN:BC (40:55). Every subsequent version adds 16 columns:
# four (Human V2, LLM Vn, Agree?) triplets + matches, comparable, %, notes.
# V4 BD:BS, V5 BT:CI, V6 CJ:CY, ... without editing this script.


def stage_layout(version):
    if version < 4:
        raise ValueError("Generic stage layout begins at V4")
    start = V3_START + 16 * (version - 3)
    return start, start + 12, start + 13, start + 14, start + 15


def stage_headers(version):
    start, matches, comparable, agreement, notes = stage_layout(version)
    result = {}
    for idx, (label, *_rest) in enumerate(CRITERIA):
        col = start + idx * 3
        result[col] = f"{label} Human V2 (reused for V{version})"
        result[col + 1] = f"{label} LLM V{version}"
        result[col + 2] = f"{label} V{version} Agree?"
    result.update({
        matches: f"V{version} Matches",
        comparable: f"V{version} Comparable",
        agreement: f"V{version} Agreement %",
        notes: f"V{version} Disagreement Notes",
    })
    return result


def stage_sheet(version):
    return f"LLM Evaluation V{version}"


def check_stage_headers(ws, version, *, allow_empty):
    for col, expected in (expected_headers() if version == 3 else stage_headers(version)).items():
        actual = ws.cell(1, col).value
        if actual != expected and not (allow_empty and actual is None):
            raise ValueError(
                f"{COMPARISON}!{get_column_letter(col)}1 is {actual!r}; "
                f"expected {expected!r}. Refusing to overwrite columns."
            )
        if allow_empty and actual is None:
            # An unlabeled occupied block must never be claimed as a version.
            for row in range(2, ws.max_row + 1):
                if ws.cell(row, col).value is not None:
                    raise ValueError(
                        f"{COMPARISON}!{get_column_letter(col)}{row} contains "
                        "unlabeled data. Refusing to overwrite it."
                    )


def ensure_stage_headers(ws, version):
    check_stage_headers(ws, version, allow_empty=True)
    _, _, _, _, notes = stage_layout(version)
    for col, title in stage_headers(version).items():
        ws.cell(1, col).value = title
        style_header(ws.cell(1, col))
        ws.column_dimensions[get_column_letter(col)].width = 55 if col == notes else 27
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 0, 48)


def ensure_stage_sheet(wb, version):
    sheet_name = stage_sheet(version)
    prior_sheet = stage_sheet(version - 1)
    headers = (
        "Question ID", "Category", "Generation Status", "Correctness",
        "Completeness", "Citation Support", "Abstention/Limits",
        "Correctness Reason", "Completeness Reason", "Citation Reason",
        "Abstention Reason", "Raw Judge Record",
    )
    if sheet_name in wb:
        ws = wb[sheet_name]
        for col, title in enumerate(headers, 1):
            if ws.cell(1, col).value != title:
                raise ValueError(f"Unexpected {sheet_name} header in column {col}")
        return ws
    ws = wb.create_sheet(sheet_name, wb.sheetnames.index(prior_sheet) + 1)
    ws.append(headers)
    for cell in ws[1]:
        style_header(cell)
    ws.row_dimensions[1].height = 38
    ws.freeze_panes = "D2"
    for col, width in {
        "A": 19, "B": 24, "C": 23, "D": 18, "E": 18,
        "F": 18, "G": 22, "H": 55, "I": 55, "J": 55,
        "K": 55, "L": 75,
    }.items():
        ws.column_dimensions[col].width = width
    return ws


def extend_stage_comparison(wb, version, human_rows, judge_rows):
    ws = wb[COMPARISON]
    ensure_stage_headers(ws, version)
    start, matches, comparable, agreement, notes = stage_layout(version)
    rows = comparison_rows(ws, wb)
    sheet_name = stage_sheet(version)
    new_count = 0
    for qid, human_row in human_rows.items():
        row = rows.get(qid)
        if row is None:
            row = max(ws.max_row + 1, 2)
            rows[qid] = row
            new_count += 1
            for col in range(1, 8):
                ref = f"'{HUMAN}'!{get_column_letter(col)}{human_row}"
                ws.cell(row, col).value = f'=IF({ref}="","",{ref})'
            # New question, no historical LLM grades; Human V2 scores remain
            # available in the existing Human V2 comparison columns.
            for (_label, _key, hcol, *_), old_col in zip(
                CRITERIA, ("K", "Q", "W", "AC")
            ):
                ws[f"{old_col}{row}"] = numeric_or_na_link(HUMAN, hcol, human_row)
        for idx, (_label, _key, hcol, jcol, _reason) in enumerate(CRITERIA):
            base = start + idx * 3
            h = f"{get_column_letter(base)}{row}"
            j = f"{get_column_letter(base + 1)}{row}"
            ws.cell(row, base).value = numeric_or_na_link(HUMAN, hcol, human_row)
            ws.cell(row, base + 1).value = numeric_or_na_link(sheet_name, jcol, judge_rows.get(qid))
            ws.cell(row, base + 2).value = formula_for_match(h, j)
        match_cells = [f"{get_column_letter(start + i * 3 + 2)}{row}" for i in range(4)]
        ws.cell(row, matches).value = "=" + "+".join(
            f'COUNTIF({m},"Match")' for m in match_cells
        )
        ws.cell(row, comparable).value = "=" + "+".join(
            f'COUNTIF({m},"Match")+COUNTIF({m},"Disagree")' for m in match_cells
        )
        a = f"{get_column_letter(matches)}{row}"
        c = f"{get_column_letter(comparable)}{row}"
        ws.cell(row, agreement).value = f'=IF({c}=0,"",{a}/{c})'
        ws.cell(row, agreement).number_format = "0.0%"
        # Notes are human-authored and deliberately never populated/overwritten.
    ws.auto_filter.ref = f"A1:{get_column_letter(notes)}{max(ws.max_row, 2)}"
    return new_count


def stage_summary_row(ws, version):
    label = f"V{version} (LLM vs Human V2 reused)"
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == label:
            return row
    # Keep all existing V1/V2/V3 summary sections intact.
    row = max(ws.max_row + 2, 28)
    for col, title in enumerate(
        ("Version", "Exact Matches", "Comparable Scores", "Exact Agreement %"), 1
    ):
        ws.cell(row - 1, col).value = title
        style_header(ws.cell(row - 1, col))
    return row


def extend_stage_summary(wb, version):
    ws = wb[SUMMARY]
    judge_name = stage_sheet(version)
    start, matches, comparable, agreement, _notes = stage_layout(version)
    last_comp = max(wb[COMPARISON].max_row, 2)
    last_judge = max(wb[judge_name].max_row, 2)
    avg_col = find_summary_cell(ws, 1, f"LLM V{version} Avg Rubric Score")
    count_col = find_summary_cell(ws, 1, f"LLM V{version} n")
    for row, (_label, _key, _hcol, jcol, _rcol) in enumerate(CRITERIA, 2):
        ref = f"'{judge_name}'!{jcol}2:{jcol}{last_judge}"
        ws.cell(row, avg_col).value = f'=IF(COUNT({ref})=0,"",AVERAGE({ref})/2)'
        ws.cell(row, avg_col).number_format = "0.0%"
        ws.cell(row, count_col).value = f'=COUNT({ref})'
    row = stage_summary_row(ws, version)
    label = f"V{version} (LLM vs Human V2 reused)"
    if ws.cell(row, 1).value not in (None, label):
        raise ValueError(f"Summary row {row} is occupied")
    ws.cell(row, 1).value = label
    ws.cell(row, 2).value = f"=SUM('{COMPARISON}'!{get_column_letter(matches)}2:{get_column_letter(matches)}{last_comp})"
    ws.cell(row, 3).value = f"=SUM('{COMPARISON}'!{get_column_letter(comparable)}2:{get_column_letter(comparable)}{last_comp})"
    ws.cell(row, 4).value = f'=IF(C{row}=0,"",B{row}/C{row})'
    ws.cell(row, 4).number_format = "0.0%"
    for cell in ws[row][:4]:
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        cell.font = Font(bold=True)

    mae_header_row = next((r for r in range(1, ws.max_row + 1)
                           if ws.cell(r, 1).value == "Criterion"
                           and ws.cell(r, 2).value == "V1 MAE (0–2)"), None)
    if mae_header_row is not None:
        title = f"V{version} MAE (0–2)"
        mae_col = next((c for c in range(1, ws.max_column + 1)
                        if ws.cell(mae_header_row, c).value == title), None)
        if mae_col is None:
            mae_col = max((c for c in range(1, ws.max_column + 1)
                           if ws.cell(mae_header_row, c).value is not None), default=1) + 1
            while any(ws.cell(mae_header_row + i, mae_col).value is not None for i in range(5)):
                mae_col += 1
        ws.cell(mae_header_row, mae_col).value = title
        style_header(ws.cell(mae_header_row, mae_col))
        ws.column_dimensions[get_column_letter(mae_col)].width = 22
        for idx, (_label, _key, _hcol, _jcol, _reason) in enumerate(CRITERIA):
            row = mae_header_row + idx + 1
            hc = f"'{COMPARISON}'!{get_column_letter(start + idx * 3)}2:{get_column_letter(start + idx * 3)}{last_comp}"
            jc = f"'{COMPARISON}'!{get_column_letter(start + idx * 3 + 1)}2:{get_column_letter(start + idx * 3 + 1)}{last_comp}"
            ws.cell(row, mae_col).value = (
                f'=IFERROR(SUMPRODUCT(IFERROR(ABS({hc}-{jc}),0))/'
                f'SUMPRODUCT(--ISNUMBER({hc}),--ISNUMBER({jc})),"")'
            )
            ws.cell(row, mae_col).number_format = "0.00"
    else:
        row0 = ws.max_row + 2
        ws.cell(row0, 1).value = f"V{version} MAE vs Human V2 (0–2, numeric pairs only)"
        style_header(ws.cell(row0, 1))
        for idx, (label, _key, _hcol, _jcol, _reason) in enumerate(CRITERIA):
            row = row0 + 1 + idx
            hc = f"'{COMPARISON}'!{get_column_letter(start + idx * 3)}2:{get_column_letter(start + idx * 3)}{last_comp}"
            jc = f"'{COMPARISON}'!{get_column_letter(start + idx * 3 + 1)}2:{get_column_letter(start + idx * 3 + 1)}{last_comp}"
            ws.cell(row, 1).value = label
            ws.cell(row, 2).value = (
                f'=IFERROR(SUMPRODUCT(IFERROR(ABS({hc}-{jc}),0))/'
                f'SUMPRODUCT(--ISNUMBER({hc}),--ISNUMBER({jc})),"")'
            )
            ws.cell(row, 2).number_format = "0.00"


def run_stage(args, version):
    src, dst = Path(args.workbook), Path(args.output)
    if not src.is_file():
        raise FileNotFoundError(f"Input workbook not found: {src}")
    if src.resolve() == dst.resolve():
        raise ValueError("--output must differ from --workbook")
    if dst.exists() and not args.dry_run:
        raise FileExistsError(f"Output exists: {dst}. Choose a new output filename.")
    evaluations = load_index(args.evaluation, "evaluation", required=True)
    validate_evaluations(evaluations)
    answers = load_index(args.answers, "answers")
    questions = load_index(args.questions, "questions")
    wb = load_workbook(src)
    for sheet in (HUMAN, "LLM Evaluation", "LLM Evaluation V2", COMPARISON, SUMMARY, stage_sheet(version - 1)):
        if sheet not in wb:
            raise ValueError(f"Missing {sheet}; V{version} requires V{version-1} first")
    check_stage_headers(wb[COMPARISON], version - 1, allow_empty=False)
    check_stage_headers(wb[COMPARISON], version, allow_empty=True)
    human_rows = id_rows(wb[HUMAN])
    comp_rows = comparison_rows(wb[COMPARISON], wb)
    if not human_rows or not comp_rows:
        raise ValueError("Expected existing human and comparison question rows")
    judge_name = stage_sheet(version)
    judge_rows = id_rows(wb[judge_name]) if judge_name in wb else {}
    if judge_name in wb:
        ensure_stage_sheet(wb, version)
    new_questions = [qid for qid in evaluations if qid not in human_rows]
    for qid in new_questions:
        m = metadata(qid, answers, questions, evaluations[qid])
        if not m[1]:
            raise ValueError(f"New question {qid} lacks question metadata; supply --answers/--questions")
    if judge_name in wb:
        judge = wb[judge_name]
        for qid, record in evaluations.items():
            if qid not in judge_rows:
                continue
            row = judge_rows[qid]
            expected = {}
            for idx, (_label, key, *_rest) in enumerate(CRITERIA):
                expected[4 + idx] = "N/A" if record["scores"][key] is None else record["scores"][key]
                expected[8 + idx] = record["scores"]["reasons"][key]
            expected[12] = json.dumps(record, ensure_ascii=False)
            for col, value in expected.items():
                old = judge.cell(row, col).value
                if old not in (None, "", value):
                    raise ValueError(f"Conflicting V{version} score for {qid} at {get_column_letter(col)}{row}")
    new_judge = [qid for qid in evaluations if qid not in judge_rows]
    new_comp = [qid for qid in evaluations if qid not in comp_rows]
    start, matches, comparable, agreement, notes = stage_layout(version)
    if args.dry_run:
        print(f"DRY RUN: V{version} records={len(evaluations)}, new judge rows={len(new_judge)}, "
              f"new questions={len(new_questions)}, new comparison rows={len(new_comp)}")
        print(f"Would append: {judge_name}; V{version} comparison "
              f"{get_column_letter(start)}:{get_column_letter(notes)}; V{version} metrics in Updated Summary")
        print("No files written; historical versions and notes unchanged.")
        return
    human = wb[HUMAN]
    judge = ensure_stage_sheet(wb, version)
    next_h = human.max_row + 1
    next_j = judge.max_row + 1
    for qid, record in evaluations.items():
        meta = metadata(qid, answers, questions, record)
        if qid not in human_rows:
            hr = next_h
            next_h += 1
            human_rows[qid] = hr
            values = (qid, meta[0], meta[1], meta[2], as_text(meta[3]),
                      meta[4], meta[5], as_text(meta[6]))
            for col, value in enumerate(values, 1):
                human.cell(hr, col).value = value
            human.cell(hr, 17).value = "New question — Human V2 review pending"
        if qid not in judge_rows:
            judge_rows[qid] = next_j
            next_j += 1
        jr = judge_rows[qid]
        for col, value in enumerate((qid, meta[0], meta[5]), 1):
            safe_fill(judge, jr, col, value, qid)
        for idx, (_label, key, *_rest) in enumerate(CRITERIA):
            score = record["scores"][key]
            safe_fill(judge, jr, 4 + idx, "N/A" if score is None else score, qid)
            safe_fill(judge, jr, 8 + idx, record["scores"]["reasons"][key], qid)
        safe_fill(judge, jr, 12, json.dumps(record, ensure_ascii=False), qid)
    judge.auto_filter.ref = f"A1:L{judge.max_row}"
    extend_stage_comparison(wb, version, human_rows, judge_rows)
    extend_stage_summary(wb, version)
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    print(f"Saved: {dst}")
    print(f"Imported V{version} records={len(evaluations)}, new judge rows={len(new_judge)}, "
          f"new questions={len(new_questions)}, new comparison rows={len(new_comp)}")
    print("Historical versions, scores, formulas, and disagreement notes preserved.")


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--workbook", required=True)
    p.add_argument("--evaluation", required=True)
    p.add_argument("--version", default="v3", help="v3, v4, v5, ... (future versions require no code changes)")
    p.add_argument("--answers", help="Required for metadata when importing unseen questions")
    p.add_argument("--questions", help="Required for metadata when importing unseen questions")
    p.add_argument("--output", required=True, help="New output workbook filename")
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()
    m = re.fullmatch(r"v([3-9]|[1-9][0-9]+)", args.version.lower().strip())
    if not m:
        raise ValueError("--version must be v3, v4, v5, ...")
    version = int(m.group(1))
    if version == 3:
        run(args)  # Retain the original, tested V3 import behavior.
    else:
        run_stage(args, version)


if __name__ == "__main__":
    main()
